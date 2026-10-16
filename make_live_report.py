#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_live_report.py — 直播逐字稿 → AI 重點摘要 Word 報告
========================================================
讀取直播逐字稿 md（.link.txt 內有 YouTube 連結），用 LLM 整理成
結構化報告（一、影片資訊  二、摘要  三、主要重點  四、重點與提醒  五、結語），
以 python-docx 產出 Word 檔存到逐字稿同層的「直播報告」資料夾。

- 20VTA 波段王子直播：沿用原本財經版型（大盤／技術分析／個股解析）。
- 其他頻道（如三師爸）：依來源自動改用通用版型，講師取自 yt-dlp 上傳者。

用法:
  python make_live_report.py <md路徑> [--out 目錄]
  python make_live_report.py --all        # 掃描「直播」目錄，補做缺報告的逐字稿
"""
import argparse
import json
import os
import re
import sys
import urllib.request
import urllib.error
from datetime import datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

PROJECT_ROOT = Path(__file__).parent
LIVE_DIR = PROJECT_ROOT / "knowledge-base" / "20VTA 頻道" / "直播"
REPORT_DIR = LIVE_DIR / "直播報告"
GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"
MAX_CHARS = 20000  # 摘要輸入上限（逐字稿截斷）


def load_gemini_key() -> str:
    env = os.environ.get("GEMINI_API_KEY")
    if env:
        return env.strip()
    key_file = Path.home() / ".gemini_api_key"
    if key_file.exists():
        return key_file.read_text(encoding="utf-8").strip()
    sys.exit("[ERR] 找不到 Gemini API Key")

def safe_filename(name):
    return re.sub(r'[\\/:*?"<>|]', "_", name)

def read_link(md_path: Path) -> str:
    link = md_path.with_suffix(".link.txt")
    if link.exists():
        return link.read_text(encoding="utf-8", errors="replace").strip()
    return ""

def video_meta(video_id: str) -> dict:
    """用 yt-dlp 快速取得標題與上傳日期"""
    try:
        import yt_dlp
        with yt_dlp.YoutubeDL({"quiet": True, "no_warnings": True,
                               "extract_flat": True, "skip_download": True}) as ydl:
            info = ydl.extract_info(f"https://www.youtube.com/watch?v={video_id}", download=False)
        return {
            "title": info.get("title", ""),
            "upload_date": info.get("upload_date", ""),  # YYYYMMDD
            "uploader": info.get("uploader", "廖崧沂"),
        }
    except Exception:
        return {"title": "", "upload_date": "", "uploader": "廖崧沂"}

def fmt_date(ymd: str) -> str:
    if len(ymd) != 8:
        return ymd
    try:
        dt = datetime.strptime(ymd, "%Y%m%d")
        wk = "一二三四五六日"[dt.weekday()]
        return f"{dt.year}年{dt.month}月{dt.day}日（週{wk}）"
    except Exception:
        return ymd

def ai_summarize(md_path: Path, video_id: str, url: str, api_key: str) -> str:
    raw = md_path.read_text(encoding="utf-8", errors="replace")
    # 去掉 md 表頭（第一列 # 標題）以節省字元
    body = raw.strip()
    if body.startswith("#"):
        parts = body.split("\n", 1)
        body = parts[1] if len(parts) > 1 else body
    if len(body) > MAX_CHARS:
        keep = body[: MAX_CHARS // 2] + "\n…（中段省略）…\n" + body[-MAX_CHARS // 2:]
        body = keep

    meta = video_meta(video_id)
    upload_date = fmt_date(meta.get("upload_date", ""))
    title = meta.get("title", md_path.stem)
    # 依來源判斷頻道（20VTA 直播報告原本即為財經直播版型，維持不變）
    is_20vta = "20VTA" in str(md_path)
    uploader = meta.get("uploader") or ("廖崧沂" if is_20vta else "三師爸Sense Bar")
    lecturer = "廖崧沂（波段王子、點靈）" if is_20vta else uploader
    source_label = "20VTA 波段王子 廖崧沂" if is_20vta else uploader

    system = ("你是專業的財經直播內容編輯，擅長將口語逐字稿整理成結構清晰的行事重點報告。"
              if is_20vta else
              "你是專業的影片／直播內容編輯，擅長將口語逐字稿整理成結構清晰、重點完整的要點報告。")

    if is_20vta:
        section3 = """## 三、主要重點與操作策略
（依內容分小節，例如「（一）大盤解析」（二）技術分析（三）操作系統（四）個股解析（五）操作與買入建議（六）課程資訊”；每個重點用「重點1：…」開頭再附「說明」；個股解析逐支條列，含股票代號與名稱）"""
        extra_rule = "- 股票代號、名稱、價位以逐字稿為準；逐字稿未提的段落（如課程資訊）可省略。\n"
    else:
        section3 = """## 三、主要重點
（依內容分小節，例如「（一）核心觀念」（二）工具／模型實測（三）操作步驟與示範（四）注意事項與坑（五）資源與資訊來源；每個重點用「重點1：…」開頭再附「說明」；若有具體步驟、設定或代碼，務必如實條列）"""
        extra_rule = "- 依逐字稿實際主題安排小節，不要套用與內容不符的章節（如財經盤勢分析）。\n"

    user = f"""請根據以下「{source_label}」直播逐字稿，整理成繁體中文報告，嚴格依下列 Markdown 結構：

# {title}

## 一、影片資訊
- 直播時間：{upload_date}
- 講師：{lecturer}
- 主題：{title}
- YouTube：{url if url else 'https://www.youtube.com/watch?v=' + video_id}

## 二、摘要
（以 150~250 字總結直播核心觀點與整體論述）

{section3}

## 四、重點與提醒事項
（條列 4~6 個最重要的提醒）

## 五、結語
（2~3 句總結，呼應開頭）

規則：
- 只依逐字稿內容整理，不可自行杜撰股票代號、價格或事件。
{extra_rule}- 刪除口語贅詞、無意義重複與主持人插話。
- 全部使用繁體中文，避免加入免責聲明以外的評語。

以下為逐字稿：
=====
{body}
=====
"""
    from google import genai
    from google.genai import types
    client = genai.Client(api_key=api_key)
    resp = client.models.generate_content(
        model="gemini-3.5-flash-lite",
        contents=f"{system}\n\n---\n\n{user}",
        config=types.GenerateContentConfig(temperature=0.3),
    )
    return (resp.text or "").strip()


def build_docx(markdown: str, out_path: Path, title: str):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "新細明體"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.15

    def add_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "新細明體"
        run.font.size = Pt(14 if level == 1 else 12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    first_line = True
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        # 標題行
        if line.startswith("# "):
            add_heading(line[2:].strip(), 1)
            first_line = False
            continue
        if line.startswith("## "):
            add_heading(line[3:].strip(), 2)
            first_line = False
            continue
        if line.startswith("### "):
            add_heading(line[4:].strip(), 2)
            first_line = False
            continue
        # 條列
        m = re.match(r'^\s*[-*•]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.space_after = Pt(2)
            p.add_run("• " + m.group(1).strip()).font.size = Pt(12)
            first_line = False
            continue
        # 一般段落
        p = doc.add_paragraph(line.strip())
        p.paragraph_format.space_after = Pt(4)
        first_line = False

    doc.save(str(out_path))
    return out_path


def make_report(md_path: Path, out_dir: Path | None = None) -> Path | None:
    md_path = Path(md_path)
    if not md_path.exists():
        print(f"[ERR] 找不到逐字稿：{md_path}")
        return None
    # 未指定輸出目錄時，報告放在逐字稿同層的「直播報告」資料夾（各頻道各自一份）
    out_dir = Path(out_dir) if out_dir else md_path.parent / "直播報告"
    out_dir.mkdir(parents=True, exist_ok=True)
    title = md_path.stem
    out_path = out_dir / f"{safe_filename(title)}.docx"
    if out_path.exists() or Path(str(out_path).replace(".docx", ".doc")).exists():
        print(f"[SKIP] 已存在：{out_path.name}")
        return out_path

    url = read_link(md_path)
    video_id = ""
    if url:
        m = re.search(r"[?&]v=([\w-]{6,})", url)
        if m:
            video_id = m.group(1)

    print(f"[製作] {title}")
    print(f"  URL: {url or '（無 link.txt）'} video_id: {video_id or '—'}")

    api_key = load_gemini_key()
    print("  [AI] Gemini 摘要中（約 1~3 分鐘）...")
    md_out = ai_summarize(md_path, video_id, url, api_key)
    print("  [AI] 摘要完成")

    build_docx(md_out, out_path, title)
    print(f"  [OK] {out_path}")
    return out_path


def scan_all_missing():
    if not LIVE_DIR.exists():
        print(f"[ERR] 找不到直播目錄：{LIVE_DIR}")
        return
    missing = []
    for md in sorted(LIVE_DIR.glob("*.md")):
        out = md.parent / "直播報告" / f"{safe_filename(md.stem)}.docx"
        if not out.exists():
            missing.append(md)
    print(f"逐字稿 {len(list(LIVE_DIR.glob('*.md')))} 支，缺報告 {len(missing)} 支")
    return missing


def main():
    ap = argparse.ArgumentParser(description="20VTA 直播逐字稿 → AI 摘要 Word 報告")
    ap.add_argument("md", nargs="?", type=Path, help="逐字稿 md 路徑")
    ap.add_argument("--out", type=Path, default=None, help="輸出目錄（預設：逐字稿同層的「直播報告」資料夾）")
    ap.add_argument("--all", action="store_true", help="補做所有缺報告的逐字稿")
    args = ap.parse_args()

    if args.all:
        missing = scan_all_missing()
        if not missing:
            print("沒有缺報告的逐字稿")
            return
        for md in missing:
            try:
                make_report(md, args.out)
            except Exception as e:
                print(f"[ERR] {md.name}: {e}")
        return

    if not args.md:
        ap.print_help()
        return
    make_report(args.md, args.out)


if __name__ == "__main__":
    main()